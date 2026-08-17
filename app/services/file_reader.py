from pathlib import Path

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


class FileReader:

    @staticmethod
    def read(path: str) -> str:
        ext = Path(path).suffix.lower()

        if ext == ".txt":
            return Path(path).read_text(
                encoding="utf-8",
                errors="ignore",
            )

        if ext == ".pdf":
            reader = PdfReader(path)

            return "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        if ext == ".docx":
            doc = Document(path)

            return "\n".join(
                paragraph.text
                for paragraph in doc.paragraphs
            )

        if ext == ".xlsx":
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
            )

            lines = []

            for worksheet in workbook.worksheets:
                lines.append(
                    f"Sheet: {worksheet.title}"
                )

                for row in worksheet.iter_rows(values_only=True):
                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:
                        lines.append(" | ".join(values))

            workbook.close()

            return "\n".join(lines)

        raise ValueError(
            f"Unsupported file type: {ext}"
        )
