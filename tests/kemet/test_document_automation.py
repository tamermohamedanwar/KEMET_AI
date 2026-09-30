from openpyxl import load_workbook

from app.services.document_automation import (
    FIELDS,
    export_csv,
    export_json,
    export_xlsx,
    extract_invoice_fields,
    process_files,
    process_records,
    validate_record,
)


def test_invoice_extraction_and_validation():
    text = """Invoice Number: INV-100\nDate: 2026-09-23\nSupplier: Acme Supplies\nCustomer: Demo Store\nCurrency: EGP\nSubtotal: 1000\nTax: 140\nTotal: 1140\nPayment Status: unpaid"""
    record = extract_invoice_fields(text, "/tmp/invoice.pdf", "a" * 64)
    checked = validate_record(record)
    assert record["invoice_number"] == "INV-100"
    assert record["document_date"] == "2026-09-23"
    assert checked["record"]["validation_status"] == "VALID"
    assert checked["errors"] == []
    assert record["field_sources"]["total"]["source_line"] == 8


def test_missing_fields_need_review():
    row = extract_invoice_fields("Invoice Number: INV-1\nSupplier: A\nTotal: 10", "a.txt", "a" * 64)
    result = process_records([row])
    assert result["summary"]["needs_review"] == 1
    assert any("missing:document_date" in e["errors"] for e in result["errors"])


def test_total_mismatch_needs_review():
    text = "Invoice Number: INV-2\nDate: 2026-09-23\nSupplier: A\nCurrency: EGP\nSubtotal: 100\nTax: 20\nTotal: 150"
    record = extract_invoice_fields(text, "b.txt", "b" * 64)
    checked = validate_record(record)
    assert checked["record"]["validation_status"] == "NEEDS_REVIEW"
    assert "total_mismatch" in checked["errors"]


def test_duplicate_source_hash_is_reported():
    a = extract_invoice_fields("Invoice Number: INV-1\nDate: 2026-09-23\nSupplier: A\nCurrency: EGP\nTotal: 10", "a.txt", "samehash")
    b = dict(a)
    b["source_file"] = "b.txt"
    result = process_records([a, b])
    assert result["summary"]["duplicates_detected"] == 1
    assert result["errors"][-1]["errors"] == ["duplicate_source_hash"]


def test_exports(tmp_path):
    row = extract_invoice_fields("Invoice Number: INV-3\nDate: 2026-09-23\nSupplier: A\nCurrency: EGP\nTotal: 10", "c.txt", "c" * 64)
    checked = validate_record(row)["record"]
    csv_path, xlsx_path, json_path = tmp_path / "clean.csv", tmp_path / "clean.xlsx", tmp_path / "report.json"
    export_csv([checked], str(csv_path))
    export_xlsx([checked], str(xlsx_path))
    export_json({"records": [checked]}, str(json_path))
    assert csv_path.read_text(encoding="utf-8-sig").splitlines()[0].startswith("document_id")
    workbook = load_workbook(xlsx_path)
    assert set(workbook.sheetnames) == {"Data", "Review", "Summary"}
    assert "INV-3" in json_path.read_text(encoding="utf-8")


def test_batch_file_pipeline_and_evidence(tmp_path):
    source = tmp_path / "invoice.txt"
    source.write_text("Invoice Number: INV-4\nDate: 2026-09-23\nSupplier: Demo\nCurrency: EGP\nTotal: 50", encoding="utf-8")
    result = process_files([str(source)])
    assert result["summary"]["records_extracted"] == 1
    assert len(result["evidence"]["input_sha256"]) == 1
    assert result["records"][0]["invoice_number"] == "INV-4"


def test_v3_quality_scorecard_and_exceptions():
    payload = {
        "records": [{"validation_status": "VALID"}],
        "errors": [{"source_file": "bad.pdf", "errors": ["total_mismatch"]}],
        "failures": [],
        "summary": {"files_received": 2, "records_extracted": 1, "valid_records": 1, "needs_review": 0, "duplicates_detected": 0, "processing_errors": 0},
        "evidence": {"input_sha256": ["abc"]},
    }
    from app.services.document_automation import build_quality_scorecard, build_exception_center, build_delivery_manifest
    score = build_quality_scorecard(payload)
    assert score["overall_quality_score"] == 75.0
    assert build_exception_center(payload)[0]["severity"] == "HIGH"
    manifest = build_delivery_manifest(payload, {"required_columns": ["invoice_number", "total"]})
    assert manifest["delivery_version"] == "V3"
    assert manifest["rules"]["required_columns"][0] == "invoice_number"


def test_v3_delivery_package(tmp_path):
    from app.services.document_automation import export_delivery_package
    row = extract_invoice_fields("Invoice Number: INV-5\nDate: 2026-09-23\nSupplier: A\nCurrency: EGP\nTotal: 10", "e.txt", "e" * 64)
    payload = process_records([row])
    payload["failures"] = []
    payload["summary"]["processing_errors"] = 0
    payload["evidence"] = {"input_sha256": ["e" * 64]}
    files = export_delivery_package(payload, str(tmp_path))
    assert set(files) == {"clean.csv", "clean.xlsx", "validation_report.json", "delivery_manifest.json"}


def test_v4_classification_and_readiness():
    from app.services.document_automation import classify_document_type, process_files_v4
    assert classify_document_type("Invoice Number: X", "x.txt")["classification_status"] == "SUPPORTED"
    assert classify_document_type("meeting notes only", "notes.txt")["classification_status"] == "UNSUPPORTED"
    result = process_files_v4([])
    assert result["readiness_gate"]["status"] == "BLOCKED"


def test_v4_client_rules_and_risk(tmp_path):
    from app.services.document_automation import process_files_v4
    source = tmp_path / "invoice.txt"
    source.write_text("Invoice Number: V4-1\nDate: 2026-09-23\nSupplier: Demo\nCurrency: USD\nSubtotal: 100\nTax: 20\nTotal: 120", encoding="utf-8")
    result = process_files_v4([str(source)], {"allowed_currencies": ["EGP"], "min_total": 200})
    codes = {r["code"] for r in result["risk_center"]}
    assert "client_currency_policy" in codes
    assert "client_min_total" in codes
    assert result["readiness_gate"]["status"] == "REVIEW_REQUIRED"


def test_v5_business_insights_and_action_plan():
    from app.services.document_automation import build_business_insights, build_action_plan
    records = [
        {"supplier_name": "A", "customer_name": "Store", "currency": "EGP", "total": "100", "payment_status": "unpaid"},
        {"supplier_name": "A", "customer_name": "Store", "currency": "EGP", "total": "50", "payment_status": "paid"},
    ]
    risks = [{"code": "duplicate_source_hash", "severity": "HIGH"}]
    insights = build_business_insights(records, risks)
    assert insights["total_value"] == 150.0
    assert insights["unpaid_or_pending_count"] == 1
    assert insights["duplicate_count"] == 1
    actions = build_action_plan(insights, risks)
    assert actions[0]["priority"] == "P1"


def test_v4_duplicate_and_evidence_package(tmp_path):
    from app.services.document_automation import export_delivery_package_v4, process_files_v4
    a = tmp_path / "a.txt"
    b = tmp_path / "b.txt"
    text = "Invoice Number: DUP-1\nDate: 2026-09-23\nSupplier: Demo\nCurrency: EGP\nSubtotal: 100\nTax: 20\nTotal: 120"
    a.write_text(text, encoding="utf-8")
    b.write_text(text, encoding="utf-8")
    result = process_files_v4([str(a), str(b)])
    assert result["summary"]["duplicates_detected"] == 1
    assert result["field_evidence"]
    assert result["quality"]["overall_quality_score"] < 100
    files = export_delivery_package_v4(result, str(tmp_path / "delivery"))
    assert set(files) == {"clean.csv", "clean.xlsx", "client_report.json", "intake_manifest.json", "validation_report.json"}
    workbook = load_workbook(files["clean.xlsx"])
    assert set(workbook.sheetnames) == {"Executive Summary", "Clean Data", "Business Insights", "Action Plan", "Review Queue", "Evidence", "Rules"}
